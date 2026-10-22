#!/usr/bin/env python3
"""Modelo cuantitativo del estudio de viabilidad: CCU, cohortes, economía y finanzas.

Ejecutar desde la raíz del repositorio:  python3 research/modelos/modelo.py
Genera: research/modelos/salidas/*.csv, research/graficos/*.png y research/modelo_financiero.xlsx

Convenciones:
- Todo dato de plataforma lleva fecha y fuente en research/13_fuentes_y_metodologia.md.
- Todo lo marcado como HIPÓTESIS es un supuesto de trabajo, no un dato observado.
- R$ = Robux. "Brutos" = Robux gastados por jugadores. "Ganados" = Earned Robux que recibe el creador.
"""
from __future__ import annotations

import math
import os
from dataclasses import dataclass

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDAS = os.path.join(RAIZ, "modelos", "salidas")
GRAFICOS = os.path.join(RAIZ, "graficos")
os.makedirs(SALIDAS, exist_ok=True)
os.makedirs(GRAFICOS, exist_ok=True)

# ---------------------------------------------------------------------------
# 1. Constantes de plataforma (OBSERVADAS, consultadas el 22-10-2026)
# ---------------------------------------------------------------------------
CUOTA_CREADOR = 0.70          # el creador recibe el 70 % de los Robux gastados en pases/productos
DEVEX_ESTANDAR = 0.0038       # USD por Earned Robux (desde 05-09-2025)
DEVEX_18_US = 0.0054          # USD por Earned Robux elegible de jugadores EE. UU. 18+ verificados (desde 08-06-2026)
DEVEX_MINIMO = 30_000         # Earned Robux mínimos por solicitud de DevEx
EURUSD = 1.1215               # tipo de referencia BCE, 20-10-2026 (USD por EUR)
USD_A_EUR = 1 / EURUSD
RECOMPENSA_DIARIA = 5         # R$ por usuario cualificado/día (Creator Rewards, Daily Engagement)
DIAS_MES = 30.4

# Plataforma (Q2 2026, carta a accionistas): 123 M DAU, 29.000 M horas, 1.557 M$ de bookings
PLAT_DAU = 123e6
PLAT_HORAS_TRIM = 29e9
PLAT_BOOKINGS_TRIM = 1.557e9
PLAT_DIAS_TRIM = 91
PLAT_DEVEX_SOBRE_BOOKINGS = 0.22  # Q2 2025 (último dato publicado encontrado)

# ---------------------------------------------------------------------------
# 2. Supuestos (HIPÓTESIS salvo que se indique lo contrario)
# ---------------------------------------------------------------------------
PERFILES_ENGAGEMENT = {  # sesiones/día, minutos/sesión
    "baja": (1.3, 10.0),
    "base": (1.6, 14.0),   # 22,4 min/día = mediana de playtime diario de GameAnalytics 2026 (observado)
    "alta": (2.0, 20.0),
}

ESCENARIOS_RETENCION = {  # D1, D7, D30
    "pesimista": (0.06, 0.010, 0.0030),
    "base": (0.12, 0.030, 0.0100),
    "optimista": (0.18, 0.060, 0.0250),
    "excepcional": (0.24, 0.100, 0.0450),
}
BENCHMARK_GA_MEDIANA = (0.103, 0.016, 0.005)   # GameAnalytics 2026 (observado; juegos con 1M+ MAU)
BENCHMARK_GA_TOP1 = (0.222, 0.091, None)

@dataclass(frozen=True)
class Monetizacion:
    conversion: float      # % de DAU que paga en el día
    arppu_rs: float        # R$ brutos por pagador y día
    cualificados_cr: float  # % de DAU que cuenta para Creator Rewards (Daily Engagement)
    cuota_18us: float      # % de Earned Robux de compras que liquida a la tasa 18+ EE. UU.

    @property
    def arpdau_rs(self) -> float:
        return self.conversion * self.arppu_rs


MONETIZACION = {
    "baja": Monetizacion(0.015, 40, 0.03, 0.00),
    "media": Monetizacion(0.025, 60, 0.05, 0.05),
    "alta": Monetizacion(0.038, 90, 0.06, 0.10),
}

CPP_USD = {"bajo": 0.01, "base": 0.03, "alto": 0.10}   # coste por partida en Ads Manager
OBJETIVOS_CCU = [10, 50, 100, 250, 500, 1000, 2000, 5000, 10000, 25000]
OBJETIVOS_EUR = [100, 500, 1000, 3000, 5000, 10000, 25000, 50000]

# Costes (EUR) - HIPÓTESIS a verificar con precios vigentes
COSTE_HERRAMIENTAS_MES = 120.0     # suscripciones IA de los dos socios
COSTE_CONTENIDO_MES = 100.0        # encargos de arte/sonido tras el lanzamiento (media)
COSTE_OTROS_MES = 30.0
COSTE_HORA_OPORTUNIDAD = 15.0      # EUR/h por socio (coste de oportunidad, no es caja)
HORAS_HASTA_LANZAMIENTO = 1600     # total de ambos socios (rango 1.300-1.950 h en el informe 10)
HORAS_MANTENIMIENTO_MES = 100      # total de ambos socios tras el lanzamiento
INVERSION_PREVIA = {               # caja antes del lanzamiento (EUR)
    "herramientas (8 meses)": 8 * COSTE_HERRAMIENTAS_MES,
    "modelos 3D, audio, UI e iconos": 1200.0,
    "anuncios de prueba (validación)": 150.0,
}


def reinversion(ingresos_eur: float) -> float:
    """Parte de los ingresos que se reinvierte en contenido/comunidad (HIPÓTESIS)."""
    if ingresos_eur >= 10_000:
        return 0.15 * ingresos_eur
    if ingresos_eur >= 3_000:
        return 0.10 * ingresos_eur
    return 0.0


def marketing_por_ccu(ccu: float) -> float:
    """Presupuesto mensual de marketing asociado a cada nivel de CCU (HIPÓTESIS)."""
    if ccu < 100:
        return 50.0
    if ccu < 1000:
        return 150.0
    if ccu < 5000:
        return 500.0
    return 1500.0


# ---------------------------------------------------------------------------
# 3. Funciones del modelo
# ---------------------------------------------------------------------------
def minutos_dia(perfil: tuple[float, float]) -> float:
    return perfil[0] * perfil[1]


def dau_por_ccu(perfil: tuple[float, float]) -> float:
    """CCU medio = DAU × sesiones × minutos / 1440  ->  DAU por cada CCU medio."""
    return 1440.0 / minutos_dia(perfil)


def curva_retencion(d1: float, d7: float, d30: float, dias: int = 730) -> np.ndarray:
    """Interpolación log-log entre anclas (día 1, 7, 30); cola con la pendiente 7→30."""
    r = np.empty(dias)
    r[0] = 1.0
    x1, x7, x30 = math.log(1), math.log(7), math.log(30)
    y1, y7, y30 = math.log(d1), math.log(d7), math.log(d30)
    p17 = (y7 - y1) / (x7 - x1)
    p730 = (y30 - y7) / (x30 - x7)
    for d in range(1, dias):
        x = math.log(d)
        if d <= 7:
            y = y1 + p17 * (x - x1)
        else:
            y = y7 + p730 * (x - x7)
        r[d] = math.exp(y)
    return r


def dias_vida(curva: np.ndarray, horizonte: int = 365) -> float:
    """Días activos esperados por usuario nuevo en el horizonte (incluye el día 0)."""
    return float(curva[:horizonte].sum())


def ingresos_por_dau_dia(m: Monetizacion) -> dict:
    """Earned Robux y USD por DAU y día."""
    brutos = m.arpdau_rs
    ganados_compras = brutos * CUOTA_CREADOR
    ganados_cr = m.cualificados_cr * RECOMPENSA_DIARIA
    usd = ganados_compras * ((1 - m.cuota_18us) * DEVEX_ESTANDAR + m.cuota_18us * DEVEX_18_US)
    usd += ganados_cr * DEVEX_ESTANDAR
    return {
        "brutos_rs": brutos,
        "ganados_compras_rs": ganados_compras,
        "ganados_cr_rs": ganados_cr,
        "ganados_rs": ganados_compras + ganados_cr,
        "usd": usd,
    }


def perfil_horario_semanal() -> np.ndarray:
    """Concurrencia relativa por hora de la semana (UTC), media = 1. Curva SINTÉTICA (HIPÓTESIS)."""
    local = np.array([0.35, 0.22, 0.15, 0.10, 0.08, 0.08, 0.12, 0.20, 0.30, 0.38, 0.45, 0.52,
                      0.60, 0.65, 0.72, 0.85, 0.95, 1.00, 1.00, 1.00, 0.95, 0.82, 0.65, 0.48])
    regiones = {  # peso de la audiencia, desfase UTC (horas)
        "Norteamérica": (0.40, -6), "Latinoamérica": (0.15, -4), "Europa": (0.25, 1),
        "Sudeste asiático": (0.12, 7), "Oceanía": (0.08, 10),
    }
    diario = np.zeros(24)
    for peso, desfase in regiones.values():
        for h_utc in range(24):
            diario[h_utc] += peso * local[(h_utc + desfase) % 24]
    dias = [0.92, 0.92, 0.92, 0.92, 1.00, 1.18, 1.12]  # lunes..domingo
    semana = np.concatenate([diario * f for f in dias])
    return semana / semana.mean()


# ---------------------------------------------------------------------------
# 4. Simulación de crecimiento por cohortes (diaria, 24 meses)
# ---------------------------------------------------------------------------
@dataclass
class EscenarioCrecimiento:
    nombre: str
    retencion: str
    engagement: tuple[float, float]
    monetizacion: str
    nuevos_lanzamiento: float   # usuarios nuevos/día extra en la ventana de lanzamiento
    nuevos_estables: float      # entrada orgánica sostenida (HIPÓTESIS central)
    vida_media_interes: float   # días para que la entrada orgánica se reduzca a la mitad sin novedades
    suelo_interes: float        # fracción de la entrada que se mantiene a largo plazo con LiveOps
    cada_actualizacion: int     # días entre actualizaciones
    pico_actualizacion: float   # +% de entrada tras cada actualización
    k_viral: float              # nuevos/día por cada DAU (invitaciones, enlaces, amigos)
    marketing_mes_eur: float
    marketing_lanzamiento_eur: float
    dia_shock: int | None = None
    factor_shock: float = 1.0
    amplificacion: float = 1.0      # multiplicador del algoritmo si las señales son buenas
    dia_amplificacion: int = 0
    rampa_amplificacion: int = 60
    dia_ruptura: int | None = None  # escenario excepcional: pico viral
    nuevos_ruptura: float = 0.0
    tau_ruptura: float = 25.0
    tam_usuarios: float = 25e6      # audiencia alcanzable (HIPÓTESIS) para saturación
    cpp_usd: float = CPP_USD["base"]


ESCENARIOS = [
    EscenarioCrecimiento("pesimista", "pesimista", (1.4, 11.0), "baja", 600, 140, 150, 0.30, 30, 0.40,
                         0.004, 0, 50, dia_shock=270, factor_shock=0.70),
    EscenarioCrecimiento("base", "base", (1.6, 14.0), "media", 2500, 1100, 240, 0.45, 21, 0.60,
                         0.010, 100, 250, dia_shock=300, factor_shock=0.75),
    EscenarioCrecimiento("optimista", "optimista", (1.8, 17.0), "alta", 5000, 3000, 300, 0.55, 14, 0.80,
                         0.020, 250, 500, amplificacion=2.2, dia_amplificacion=45, rampa_amplificacion=90),
    EscenarioCrecimiento("excepcional", "excepcional", (1.8, 19.0), "media", 6000, 3500, 300, 0.60, 14, 0.80,
                         0.030, 500, 500, amplificacion=2.0, dia_amplificacion=45, rampa_amplificacion=60,
                         dia_ruptura=90, nuevos_ruptura=160000, tau_ruptura=20),
]


def simular(e: EscenarioCrecimiento, dias: int = 730) -> pd.DataFrame:
    ret = curva_retencion(*ESCENARIOS_RETENCION[e.retencion], dias=dias)
    perfil = e.engagement
    mon = MONETIZACION[e.monetizacion]
    por_dau = ingresos_por_dau_dia(mon)
    nuevos = np.zeros(dias)
    dau = np.zeros(dias)
    acumulado = 0.0
    filas = []
    for t in range(dias):
        decaimiento = e.suelo_interes + (1 - e.suelo_interes) * 0.5 ** (t / e.vida_media_interes)
        lanzamiento = e.nuevos_lanzamiento * math.exp(-t / 10.0)
        desde_update = t % e.cada_actualizacion if t >= e.cada_actualizacion else None
        mult_update = 1.0 + (e.pico_actualizacion * math.exp(-desde_update / 6.0) if desde_update is not None else 0.0)
        amp = 1.0
        if e.amplificacion > 1 and t >= e.dia_amplificacion:
            avance = min(1.0, (t - e.dia_amplificacion) / e.rampa_amplificacion)
            amp = 1.0 + (e.amplificacion - 1.0) * avance
        shock = e.factor_shock if (e.dia_shock is not None and t >= e.dia_shock) else 1.0
        organico = (e.nuevos_estables * decaimiento * mult_update * amp + lanzamiento) * shock
        if e.dia_ruptura is not None and t >= e.dia_ruptura:
            organico += e.nuevos_ruptura * math.exp(-(t - e.dia_ruptura) / e.tau_ruptura)
        viral = e.k_viral * (dau[t - 1] if t > 0 else 0.0)
        presupuesto = e.marketing_mes_eur / DIAS_MES + (e.marketing_lanzamiento_eur / 14 if t < 14 else 0.0)
        pagados = presupuesto * EURUSD / e.cpp_usd
        saturacion = max(0.0, 1.0 - acumulado / e.tam_usuarios)
        nuevos[t] = (organico + viral) * saturacion + pagados
        acumulado += nuevos[t]
        # DAU(t) = suma de cohortes anteriores × retención a su edad
        dau[t] = float(np.dot(nuevos[: t + 1][::-1], ret[: t + 1]))
        ccu_medio = dau[t] * minutos_dia(perfil) / 1440.0
        filas.append({
            "dia": t, "nuevos": nuevos[t], "organicos": organico * saturacion, "virales": viral * saturacion,
            "pagados": pagados, "dau": dau[t], "ccu_medio": ccu_medio,
            "usd_dia": dau[t] * por_dau["usd"], "ganados_rs_dia": dau[t] * por_dau["ganados_rs"],
            "marketing_eur_dia": presupuesto,
        })
    df = pd.DataFrame(filas)
    df["mes"] = df["dia"] // 30.4
    return df


def resumen_mensual(e: EscenarioCrecimiento, df: pd.DataFrame, pico_ratio: float) -> pd.DataFrame:
    g = df.groupby(df["dia"].floordiv(DIAS_MES).astype(int))
    m = pd.DataFrame({
        "mes": g["dia"].min().floordiv(DIAS_MES).astype(int) + 1,
        "nuevos_dia_medio": g["nuevos"].mean(),
        "dau_medio": g["dau"].mean(),
        "ccu_medio": g["ccu_medio"].mean(),
        "ccu_pico_semanal_est": g["ccu_medio"].mean() * pico_ratio,
        "ingresos_usd": g["usd_dia"].sum(),
        "marketing_eur": g["marketing_eur_dia"].sum(),
    }).reset_index(drop=True)
    m = m[m["mes"] <= 24]
    m["ingresos_eur"] = m["ingresos_usd"] * USD_A_EUR
    fijos = COSTE_HERRAMIENTAS_MES + COSTE_CONTENIDO_MES + COSTE_OTROS_MES
    m["costes_eur"] = fijos + m["marketing_eur"] + m["ingresos_eur"].apply(reinversion)
    m["beneficio_operativo_eur"] = m["ingresos_eur"] - m["costes_eur"]
    inversion = sum(INVERSION_PREVIA.values())  # el marketing de lanzamiento ya va en marketing_eur
    m["caja_acumulada_eur"] = m["beneficio_operativo_eur"].cumsum() - inversion
    m["por_socio_eur"] = m["beneficio_operativo_eur"] / 2
    coste_op = HORAS_MANTENIMIENTO_MES * COSTE_HORA_OPORTUNIDAD
    m["resultado_con_coste_oportunidad_eur"] = m["beneficio_operativo_eur"] - coste_op
    m.insert(0, "escenario", e.nombre)
    return m


# ---------------------------------------------------------------------------
# 5. Tablas
# ---------------------------------------------------------------------------
def tabla_ccu(semana: np.ndarray) -> pd.DataFrame:
    pico = semana.max()
    pico_diario_laborable = semana[: 24 * 4].reshape(4, 24).max(axis=1).mean()
    minimo = semana.min()
    filas = []
    vida = {k: dias_vida(curva_retencion(*v)) for k, v in ESCENARIOS_RETENCION.items()}
    for ccu in OBJETIVOS_CCU:
        fila = {"ccu_medio": ccu,
                "ccu_pico_semanal": round(ccu * pico),
                "ccu_pico_laborable": round(ccu * pico_diario_laborable),
                "ccu_minimo": round(ccu * minimo)}
        for nombre, perfil in PERFILES_ENGAGEMENT.items():
            fila[f"dau_eng_{nombre}"] = round(ccu * dau_por_ccu(perfil))
        dau_base = ccu * dau_por_ccu(PERFILES_ENGAGEMENT["base"])
        for nombre, l in vida.items():
            fila[f"nuevos_dia_ret_{nombre}"] = round(dau_base / l)
        fila["coste_mes_usd_si_todo_pagado_cpp_0.03"] = round(dau_base / vida["base"] * DIAS_MES * CPP_USD["base"])
        filas.append(fila)
    return pd.DataFrame(filas)


def tabla_finanzas_por_ccu() -> pd.DataFrame:
    filas = []
    perfil = PERFILES_ENGAGEMENT["base"]
    for ccu in OBJETIVOS_CCU:
        dau = ccu * dau_por_ccu(perfil)
        for nombre, m in MONETIZACION.items():
            r = ingresos_por_dau_dia(m)
            usd_mes = dau * r["usd"] * DIAS_MES
            eur_mes = usd_mes * USD_A_EUR
            costes = (COSTE_HERRAMIENTAS_MES + COSTE_CONTENIDO_MES + COSTE_OTROS_MES
                      + marketing_por_ccu(ccu) + reinversion(eur_mes))
            beneficio = eur_mes - costes
            filas.append({
                "ccu_medio": ccu, "monetizacion": nombre, "dau": round(dau),
                "pagadores_dia": round(dau * m.conversion, 1),
                "arpdau_rs_brutos": round(m.arpdau_rs, 2),
                "rs_brutos_dia": round(dau * r["brutos_rs"]),
                "earned_rs_mes": round(dau * r["ganados_rs"] * DIAS_MES),
                "usd_dia": round(dau * r["usd"], 2),
                "usd_mes": round(usd_mes),
                "eur_mes": round(eur_mes),
                "eur_anio": round(eur_mes * 12),
                "costes_eur_mes": round(costes),
                "beneficio_operativo_eur_mes": round(beneficio),
                "por_socio_eur_mes": round(beneficio / 2),
                "beneficio_con_coste_oportunidad_eur_mes": round(beneficio - HORAS_MANTENIMIENTO_MES * COSTE_HORA_OPORTUNIDAD),
            })
    return pd.DataFrame(filas)


def eur_por_ccu_mes(m: Monetizacion, perfil=PERFILES_ENGAGEMENT["base"]) -> float:
    return dau_por_ccu(perfil) * ingresos_por_dau_dia(m)["usd"] * DIAS_MES * USD_A_EUR


def tabla_objetivos() -> pd.DataFrame:
    filas = []
    perfil = PERFILES_ENGAGEMENT["base"]
    vida_base = dias_vida(curva_retencion(*ESCENARIOS_RETENCION["base"]))
    vida_opt = dias_vida(curva_retencion(*ESCENARIOS_RETENCION["optimista"]))
    for objetivo in OBJETIVOS_EUR:
        for nombre, m in MONETIZACION.items():
            ccu = objetivo / eur_por_ccu_mes(m, perfil)
            dau = ccu * dau_por_ccu(perfil)
            filas.append({
                "objetivo_eur_mes": objetivo, "monetizacion": nombre,
                "ccu_medio_necesario": round(ccu, 1),
                "dau_necesario": round(dau),
                "pagadores_dia": round(dau * m.conversion, 1),
                "rs_brutos_mes": round(dau * m.arpdau_rs * DIAS_MES),
                "nuevos_dia_ret_base": round(dau / vida_base),
                "nuevos_dia_ret_optimista": round(dau / vida_opt),
            })
    return pd.DataFrame(filas)


def tabla_ltv() -> pd.DataFrame:
    filas = []
    for rn, anclas in ESCENARIOS_RETENCION.items():
        l = dias_vida(curva_retencion(*anclas))
        for mn, m in MONETIZACION.items():
            ltv = l * ingresos_por_dau_dia(m)["usd"]
            filas.append({"retencion": rn, "monetizacion": mn, "dias_activos_365": round(l, 2),
                          "ltv_usd_creador": round(ltv, 4),
                          "cpp_maximo_rentable_usd": round(ltv, 4),
                          "roi_cpp_0.01": round(ltv / 0.01, 2), "roi_cpp_0.03": round(ltv / 0.03, 2)})
    return pd.DataFrame(filas)


def tabla_equilibrio() -> pd.DataFrame:
    fijos = COSTE_HERRAMIENTAS_MES + COSTE_CONTENIDO_MES + COSTE_OTROS_MES + 50  # + marketing mínimo
    oportunidad = HORAS_MANTENIMIENTO_MES * COSTE_HORA_OPORTUNIDAD
    filas = []
    for nombre, m in MONETIZACION.items():
        e = eur_por_ccu_mes(m)
        filas.append({"monetizacion": nombre, "eur_por_ccu_medio_mes": round(e, 2),
                      "costes_caja_mes": fijos, "ccu_equilibrio_caja": round(fijos / e, 1),
                      "coste_oportunidad_mes": oportunidad,
                      "ccu_equilibrio_con_oportunidad": round((fijos + oportunidad) / e, 1)})
    return pd.DataFrame(filas)


# ---------------------------------------------------------------------------
# 6. Matriz de decisión de conceptos (puntuaciones de juicio experto 0-10)
# ---------------------------------------------------------------------------
CRITERIOS = [  # (clave, nombre, peso base %)
    ("captacion", "Facilidad de captación (CTR/PTR)", 10),
    ("retencion", "Potencial de retención (D1-D28)", 16),
    ("monetizacion", "Potencial de monetización", 12),
    ("viralidad", "Viralidad y co-juego", 8),
    ("originalidad", "Originalidad / diferenciación", 8),
    ("tecnica", "Viabilidad técnica (inversa a complejidad)", 10),
    ("contenido", "Contenido necesario (inverso)", 8),
    ("tiempo", "Tiempo de desarrollo (inverso)", 8),
    ("riesgo", "Riesgo de ejecución (inverso)", 6),
    ("competencia", "Competencia directa (inversa)", 8),
    ("crecimiento", "Potencial de crecimiento", 6),
]
CONCEPTOS = {
    "C1 Carreras callejeras mundo abierto + reputación": [8, 6, 7, 6, 5, 4, 4, 4, 5, 4, 8],
    "C2 Arcade de personalización / car meets": [6, 5, 8, 7, 6, 6, 5, 6, 5, 6, 6],
    "C3 PvP clasificatorio por temporadas": [5, 6, 4, 6, 6, 5, 7, 7, 4, 7, 5],
    "C4 De coche barato a hiperdeportivo (colección/tycoon)": [7, 7, 8, 4, 3, 6, 4, 6, 5, 3, 6],
    "C5 Mundo abierto social / roleplay": [7, 6, 6, 5, 2, 3, 3, 3, 4, 2, 7],
    "C6 Street racing por crews + persecuciones (híbrido)": [7, 8, 7, 8, 8, 5, 6, 6, 5, 7, 8],
}
VARIANTES_PESOS = {
    "Base": {},
    "Prioridad rentabilidad": {"monetizacion": 2.5, "captacion": 1.3},
    "Prioridad velocidad de lanzamiento": {"tiempo": 2.5, "contenido": 2.0, "tecnica": 1.8},
    "Prioridad retención": {"retencion": 2.0, "viralidad": 1.5},
    "Prioridad bajo riesgo": {"riesgo": 2.5, "competencia": 1.5, "tecnica": 1.5},
}


def matriz_decision() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    claves = [c[0] for c in CRITERIOS]
    base = np.array([c[2] for c in CRITERIOS], dtype=float)
    puntuaciones = pd.DataFrame(CONCEPTOS, index=claves).T
    resultados = {}
    for variante, mult in VARIANTES_PESOS.items():
        w = base.copy()
        for k, f in mult.items():
            w[claves.index(k)] *= f
        w = w / w.sum()
        resultados[variante] = (puntuaciones.values @ w) * 10  # escala 0-100
    tabla = pd.DataFrame(resultados, index=puntuaciones.index).round(1)
    # Monte Carlo: pesos aleatorios (Dirichlet centrada en los pesos base) + ruido ±1 en puntuaciones
    rng = np.random.default_rng(20261022)
    victorias = dict.fromkeys(puntuaciones.index, 0)
    n = 20000
    alpha = base / base.sum() * 40
    for _ in range(n):
        w = rng.dirichlet(alpha)
        ruido = rng.integers(-1, 2, size=puntuaciones.shape)
        p = np.clip(puntuaciones.values + ruido, 0, 10)
        victorias[puntuaciones.index[int(np.argmax(p @ w))]] += 1
    mc = pd.DataFrame({"prob_ganar_%": {k: round(100 * v / n, 1) for k, v in victorias.items()}})
    # Prueba de estrés: ¿y si el componente cooperativo de C6 no funciona como esperamos?
    w = base / base.sum()
    estres = puntuaciones.copy()
    c6 = [i for i in estres.index if i.startswith("C6")][0]
    for k, d in {"retencion": -2, "viralidad": -2, "originalidad": -2, "tecnica": -1, "tiempo": -1}.items():
        estres.loc[c6, k] += d
    mc["estres_C6_puntuacion"] = (estres.values @ w * 10).round(1)
    return puntuaciones, tabla, mc


# ---------------------------------------------------------------------------
# 7. Economía del juego: tiempo para conseguir coches (HIPÓTESIS de diseño)
# ---------------------------------------------------------------------------
CLASES = [  # clase, precio medio coche (créditos), ingreso/hora jugando en la clase anterior, % gastado en mejoras
    ("D", 0, None, 0.0),
    ("C", 70_000, 32_000, 0.15),
    ("B", 230_000, 58_000, 0.25),
    ("A", 1_000_000, 100_000, 0.30),
    ("S", 4_000_000, 170_000, 0.30),
]
BONO_INICIAL = 15_000          # tutorial + primeros hitos
MISIONES_DIARIAS_HORA = 8_000  # créditos/hora equivalentes de diarias para un jugador de ~40 min/día


def tabla_economia() -> pd.DataFrame:
    filas, acum, acum_doble = [], 0.0, 0.0
    for clase, precio, ingreso, mejoras in CLASES[1:]:
        objetivo = precio - (BONO_INICIAL if clase == "C" else 0)
        horas = objetivo / ((ingreso + MISIONES_DIARIAS_HORA) * (1 - mejoras))
        # el pase duplica los premios de evento, no las misiones
        horas_doble = objetivo / ((2 * ingreso + MISIONES_DIARIAS_HORA) * (1 - mejoras))
        acum += horas
        acum_doble += horas_doble
        filas.append({"clase": clase, "precio_coche": precio, "ingreso_hora_previo": ingreso,
                      "horas_tramo_gratis": round(horas, 1), "horas_acumuladas_gratis": round(acum, 1),
                      "horas_acumuladas_doble_ganancias": round(acum_doble, 1)})
    return pd.DataFrame(filas)


# ---------------------------------------------------------------------------
# 8. Sensibilidad (tornado) sobre el escenario base
# ---------------------------------------------------------------------------
def sensibilidad(base: EscenarioCrecimiento) -> pd.DataFrame:
    def metricas(e: EscenarioCrecimiento, mon: Monetizacion | None = None, ret=None) -> tuple[float, float]:
        if ret is not None:
            ESCENARIOS_RETENCION["_tmp"] = ret
            e = EscenarioCrecimiento(**{**e.__dict__, "retencion": "_tmp"})
        if mon is not None:
            MONETIZACION["_tmp"] = mon
            e = EscenarioCrecimiento(**{**e.__dict__, "monetizacion": "_tmp"})
        df = simular(e, dias=366)
        ccu12 = df[(df["dia"] >= 335) & (df["dia"] < 366)]["ccu_medio"].mean()
        ingresos = df["usd_dia"].sum() * USD_A_EUR
        ESCENARIOS_RETENCION.pop("_tmp", None)
        MONETIZACION.pop("_tmp", None)
        return ccu12, ingresos

    ccu0, ing0 = metricas(base)
    r = ESCENARIOS_RETENCION[base.retencion]
    m = MONETIZACION[base.monetizacion]
    pruebas = {
        "Entrada orgánica sostenida": lambda f: metricas(EscenarioCrecimiento(**{**base.__dict__, "nuevos_estables": base.nuevos_estables * f})),
        "Retención D1": lambda f: metricas(base, ret=(r[0] * f, r[1], r[2])),
        "Retención D7 y D30 (juntas)": lambda f: metricas(base, ret=(r[0], r[1] * f, r[2] * f)),
        "Retención D30 (cola larga)": lambda f: metricas(base, ret=(r[0], r[1], r[2] * f)),
        "Minutos por sesión": lambda f: metricas(EscenarioCrecimiento(**{**base.__dict__, "engagement": (base.engagement[0], base.engagement[1] * f)})),
        "Factor viral k": lambda f: metricas(EscenarioCrecimiento(**{**base.__dict__, "k_viral": base.k_viral * f})),
        "Vida media del interés": lambda f: metricas(EscenarioCrecimiento(**{**base.__dict__, "vida_media_interes": base.vida_media_interes * f})),
        "Presupuesto de marketing": lambda f: metricas(EscenarioCrecimiento(**{**base.__dict__, "marketing_mes_eur": base.marketing_mes_eur * f})),
        "Conversión a pagador": lambda f: metricas(base, mon=Monetizacion(m.conversion * f, m.arppu_rs, m.cualificados_cr, m.cuota_18us)),
        "ARPPU": lambda f: metricas(base, mon=Monetizacion(m.conversion, m.arppu_rs * f, m.cualificados_cr, m.cuota_18us)),
    }
    filas = []
    for nombre, fn in pruebas.items():
        bajo = fn(0.7)
        alto = fn(1.3)
        filas.append({"variable": nombre,
                      "ccu_mes12_-30%": round(bajo[0], 1), "ccu_mes12_+30%": round(alto[0], 1),
                      "ingresos_12m_eur_-30%": round(bajo[1]), "ingresos_12m_eur_+30%": round(alto[1])})
    df = pd.DataFrame(filas)
    df.attrs["ccu_base"] = ccu0
    df.attrs["ingresos_base"] = ing0
    return df


# ---------------------------------------------------------------------------
# 9. Gráficos (paleta de referencia validada: azul, naranja, aqua, amarillo)
# ---------------------------------------------------------------------------
SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100"]
SUPERFICIE, TINTA, TINTA2, TENUE, REJILLA, EJE = "#fcfcfb", "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7"


def estilo(ax, titulo: str, subtitulo: str | None = None):
    ax.set_facecolor(SUPERFICIE)
    ax.figure.set_facecolor(SUPERFICIE)
    for lado in ("top", "right"):
        ax.spines[lado].set_visible(False)
    for lado in ("left", "bottom"):
        ax.spines[lado].set_color(EJE)
    ax.tick_params(colors=TENUE, labelsize=9)
    ax.grid(True, color=REJILLA, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.xaxis.label.set_color(TINTA2)
    ax.yaxis.label.set_color(TINTA2)
    ax.set_title(titulo, loc="left", color=TINTA, fontsize=12, fontweight="bold", pad=22 if subtitulo else 10)
    if subtitulo:
        ax.text(0, 1.02, subtitulo, transform=ax.transAxes, color=TINTA2, fontsize=9)


def fmt_num(v: float) -> str:
    if v >= 1e6:
        return f"{v / 1e6:.1f} M".replace(".", ",")
    if v >= 1e3:
        return f"{v / 1e3:.1f} k".replace(".", ",")
    return f"{v:.0f}"


def guardar(fig, nombre: str):
    fig.tight_layout()
    fig.savefig(os.path.join(GRAFICOS, nombre), dpi=150, facecolor=SUPERFICIE)
    plt.close(fig)


def graficos(semana, sims, mensual, fin, objetivos, sens, matriz, competidores, economia):
    # g01 DAU necesario por CCU
    fig, ax = plt.subplots(figsize=(8, 4.8))
    for i, (nombre, perfil) in enumerate(PERFILES_ENGAGEMENT.items()):
        y = [c * dau_por_ccu(perfil) for c in OBJETIVOS_CCU]
        ax.plot(OBJETIVOS_CCU, y, color=SERIES[i], lw=2, marker="o", ms=5, mec=SUPERFICIE, mew=1.5,
                label=f"Engagement {nombre} ({minutos_dia(perfil):.0f} min/día)")
        ax.annotate(f"{fmt_num(y[-1])} DAU", (OBJETIVOS_CCU[-1], y[-1]), xytext=(6, 0), textcoords="offset points",
                    color=TINTA2, fontsize=8, va="center")
    ax.set_xscale("log")
    ax.set_yscale("log")
    ax.set_xlabel("CCU medio objetivo")
    ax.set_ylabel("DAU necesario")
    ax.legend(frameon=False, fontsize=8, labelcolor=TINTA2)
    estilo(ax, "DAU necesario para cada CCU medio", "CCU medio = DAU × min/día ÷ 1440 · escalas logarítmicas")
    ax.set_xlim(8, 60000)
    guardar(fig, "g01_dau_por_ccu.png")

    # g02 perfil semanal
    fig, ax = plt.subplots(figsize=(9, 3.6))
    ax.plot(np.arange(len(semana)), semana, color=SERIES[0], lw=2)
    ax.fill_between(np.arange(len(semana)), semana, color=SERIES[0], alpha=0.10)
    ax.axhline(1.0, color=EJE, lw=1)
    imax, imin = int(np.argmax(semana)), int(np.argmin(semana))
    ax.annotate(f"pico ×{semana[imax]:.2f}", (imax, semana[imax]), xytext=(4, 4), textcoords="offset points", color=TINTA2, fontsize=8)
    ax.annotate(f"mínimo ×{semana[imin]:.2f}", (imin, semana[imin]), xytext=(4, -12), textcoords="offset points", color=TINTA2, fontsize=8)
    ax.set_xticks(np.arange(0, 168, 24) + 12)
    ax.set_xticklabels(["lun", "mar", "mié", "jue", "vie", "sáb", "dom"])
    ax.set_ylabel("CCU relativo a la media")
    estilo(ax, "Concurrencia por hora de la semana (curva sintética, UTC)", "Hipótesis de mezcla geográfica: 55 % América, 25 % Europa, 20 % Asia-Pacífico")
    guardar(fig, "g02_perfil_semanal.png")

    # g03 retención
    fig, ax = plt.subplots(figsize=(8, 4.6))
    dias = np.arange(1, 91)
    for i, (nombre, anclas) in enumerate(ESCENARIOS_RETENCION.items()):
        curva = curva_retencion(*anclas)[1:91] * 100
        ax.plot(dias, curva, color=SERIES[i], lw=2, label=f"{nombre} (D1 {anclas[0]*100:.0f} %, D7 {anclas[1]*100:.0f} %, D30 {anclas[2]*100:.1f} %)".replace(".", ","))
    ax.scatter([1, 7, 30], [v * 100 for v in BENCHMARK_GA_MEDIANA], color=TINTA2, s=36, zorder=5, edgecolors=SUPERFICIE, linewidths=1.5,
               label="Mediana GameAnalytics 2026 (observado)")
    ax.set_yscale("log")
    ax.set_xlabel("Días desde la primera sesión")
    ax.set_ylabel("% de la cohorte activa (escala log)")
    ax.legend(frameon=False, fontsize=8, labelcolor=TINTA2)
    estilo(ax, "Curvas de retención usadas en el modelo", "Escenarios hipotéticos frente a la mediana observada de juegos con 1M+ MAU")
    guardar(fig, "g03_retencion.png")

    # g04 simulación CCU 24 meses
    fig, ax = plt.subplots(figsize=(9, 4.8))
    for i, (nombre, df) in enumerate(sims.items()):
        semanal = df["ccu_medio"].rolling(7, min_periods=1).mean()
        ax.plot(df["dia"] / DIAS_MES, semanal, color=SERIES[i], lw=2, label=nombre)
        ax.annotate(f"{nombre}: {fmt_num(semanal.iloc[-1])}", (24, semanal.iloc[-1]), xytext=(6, 0), textcoords="offset points",
                    color=TINTA2, fontsize=8, va="center")
    ax.set_yscale("log")
    ax.set_xlabel("Meses desde el lanzamiento")
    ax.set_ylabel("CCU medio (media móvil 7 días, log)")
    ax.set_xlim(0, 30)
    ax.legend(frameon=False, fontsize=8, labelcolor=TINTA2, loc="upper center", bbox_to_anchor=(0.5, -0.14), ncol=4)
    estilo(ax, "Simulación de CCU medio a 24 meses", "Escenarios hipotéticos; el excepcional no es una previsión")
    guardar(fig, "g04_simulacion_ccu.png")

    # g05 ingresos por CCU
    fig, ax = plt.subplots(figsize=(8, 4.8))
    for i, nombre in enumerate(MONETIZACION):
        sub = fin[fin["monetizacion"] == nombre]
        ax.plot(sub["ccu_medio"], sub["eur_mes"], color=SERIES[i], lw=2, marker="o", ms=5, mec=SUPERFICIE, mew=1.5,
                label=f"Monetización {nombre}")
        ultimo = sub.iloc[-1]
        ax.annotate(f"{fmt_num(ultimo['eur_mes'])} €", (ultimo["ccu_medio"], ultimo["eur_mes"]), xytext=(6, 0),
                    textcoords="offset points", color=TINTA2, fontsize=8, va="center")
    ax.set_xscale("log")
    ax.set_yscale("log")
    ax.set_xlabel("CCU medio")
    ax.set_ylabel("Ingresos netos para el estudio (€/mes, log)")
    ax.set_xlim(8, 60000)
    ax.legend(frameon=False, fontsize=8, labelcolor=TINTA2)
    estilo(ax, "Ingresos mensuales del estudio según CCU medio", "Tras la cuota de Roblox y DevEx; antes de costes e impuestos")
    guardar(fig, "g05_ingresos_por_ccu.png")

    # g06 CCU necesario por objetivo de ingresos (puntos unidos; escala log, sin barras)
    fig, ax = plt.subplots(figsize=(9, 4.8))
    x = np.arange(len(OBJETIVOS_EUR))
    for i, nombre in enumerate(MONETIZACION):
        sub = objetivos[objetivos["monetizacion"] == nombre]
        ax.plot(x, sub["ccu_medio_necesario"], color=SERIES[i], lw=2, marker="o", ms=6, mec=SUPERFICIE, mew=1.5,
                label=f"Monetización {nombre}")
        ultimo = sub["ccu_medio_necesario"].iloc[-1]
        ax.annotate(f"{fmt_num(ultimo)} CCU", (x[-1], ultimo), xytext=(8, 0), textcoords="offset points",
                    color=TINTA2, fontsize=8, va="center")
    ax.set_yscale("log")
    ax.set_xticks(x)
    ax.set_xticklabels([f"{v:,}".replace(",", ".") + " €" for v in OBJETIVOS_EUR], fontsize=8)
    ax.set_xlim(-0.3, len(x) - 0.3)
    ax.set_xlabel("Ingresos netos mensuales objetivo (antes de costes)")
    ax.set_ylabel("CCU medio necesario (log)")
    ax.legend(frameon=False, fontsize=8, labelcolor=TINTA2)
    estilo(ax, "CCU medio necesario para cada objetivo de ingresos mensuales", "Tras cuota de Roblox y DevEx; engagement base (22,4 min/día)")
    guardar(fig, "g06_ccu_por_objetivo.png")

    # g07 tornados: CCU del mes 12 e ingresos de 12 meses
    for col_bajo, col_alto, base_val, fichero, titulo, eje, fmt in (
        ("ccu_mes12_-30%", "ccu_mes12_+30%", sens.attrs["ccu_base"], "g07_tornado_ccu.png",
         "Sensibilidad del CCU del mes 12 (±30 % en cada variable)", "CCU medio en el mes 12 (escenario base)", "{:.0f}"),
        ("ingresos_12m_eur_-30%", "ingresos_12m_eur_+30%", sens.attrs["ingresos_base"], "g07b_tornado_ingresos.png",
         "Sensibilidad de los ingresos de los 12 primeros meses (±30 %)", "Ingresos netos acumulados en 12 meses (€)", "{:,.0f} €")):
        s2 = sens.copy()
        s2["rango"] = (s2[col_alto] - s2[col_bajo]).abs()
        s2 = s2[s2["rango"] > 1e-6].sort_values("rango")
        fig, ax = plt.subplots(figsize=(8, 4.6))
        y = np.arange(len(s2))
        ax.barh(y, s2[col_bajo] - base_val, left=base_val, color=SERIES[1], height=0.5, label="Variable −30 %")
        ax.barh(y, s2[col_alto] - base_val, left=base_val, color=SERIES[0], height=0.5, label="Variable +30 %")
        ax.axvline(base_val, color=TINTA2, lw=1)
        ax.set_yticks(y)
        ax.set_yticklabels(s2["variable"], fontsize=8)
        ax.set_xlabel(eje)
        ax.legend(frameon=False, fontsize=8, labelcolor=TINTA2, loc="lower right")
        estilo(ax, titulo, "Escenario base: " + fmt.format(base_val).replace(",", "."))
        guardar(fig, fichero)

    # g08 matriz de decisión (serie única, ganador resaltado)
    t = matriz.sort_values("Base")
    fig, ax = plt.subplots(figsize=(8, 4.2))
    colores = [SERIES[0] if v == t["Base"].max() else EJE for v in t["Base"]]
    ax.barh(np.arange(len(t)), t["Base"], color=colores, height=0.5)
    for i, v in enumerate(t["Base"]):
        ax.text(v + 0.8, i, f"{v:.1f}".replace(".", ","), va="center", color=TINTA2, fontsize=8)
    ax.set_yticks(np.arange(len(t)))
    cortos = {"C1": "C1 Callejeras en mundo abierto", "C2": "C2 Personalización / car meets", "C3": "C3 PvP por temporadas",
              "C4": "C4 Colección / tycoon", "C5": "C5 Social / roleplay", "C6": "C6 Crews + persecuciones (híbrido)"}
    ax.set_yticklabels([cortos[n[:2]] for n in t.index], fontsize=8)
    ax.set_xlim(0, 100)
    ax.set_xlabel("Puntuación ponderada (0-100)")
    estilo(ax, "Matriz de decisión de conceptos (pesos base)")
    guardar(fig, "g08_matriz_conceptos.png")

    # g09 competidores (instantáneas): gráfico de puntos en escala log (las barras no deben ir en log)
    c = competidores.dropna(subset=["ccu_instantanea"]).sort_values("ccu_instantanea")
    fig, ax = plt.subplots(figsize=(8, 5.4))
    ax.hlines(np.arange(len(c)), 10, c["ccu_instantanea"], color=REJILLA, lw=1)
    ax.scatter(c["ccu_instantanea"], np.arange(len(c)), color=SERIES[0], s=48, zorder=3, edgecolors=SUPERFICIE, linewidths=1.5)
    for i, v in enumerate(c["ccu_instantanea"]):
        ax.text(v * 1.15, i, fmt_num(v), va="center", color=TINTA2, fontsize=8)
    ax.set_yticks(np.arange(len(c)))
    ax.set_yticklabels(c["juego"], fontsize=8)
    ax.set_xscale("log")
    ax.set_xlim(10, 80000)
    ax.set_xlabel("Jugadores concurrentes en una instantánea (log)")
    estilo(ax, "Competidores: CCU en instantáneas de jul-oct 2026", "Rastreadores de terceros; fechas y fuentes en 02_competidores.md")
    guardar(fig, "g09_competidores_ccu.png")

    # g10 caja acumulada
    fig, ax = plt.subplots(figsize=(9, 4.8))
    for i, (nombre, m) in enumerate(mensual.items()):
        ax.plot(m["mes"], m["caja_acumulada_eur"], color=SERIES[i], lw=2, marker="o", ms=3)
        ax.annotate(f"{nombre}: {fmt_num(abs(m['caja_acumulada_eur'].iloc[-1]))} €" + (" (pérdida)" if m["caja_acumulada_eur"].iloc[-1] < 0 else ""),
                    (24, m["caja_acumulada_eur"].iloc[-1]), xytext=(6, 0), textcoords="offset points", color=TINTA2, fontsize=8, va="center")
    ax.axhline(0, color=TINTA2, lw=1)
    ax.set_yscale("symlog", linthresh=5000)
    ax.set_xlabel("Meses desde el lanzamiento")
    ax.set_ylabel("Caja acumulada (€, escala symlog)")
    ax.set_xlim(0, 32)
    estilo(ax, "Caja acumulada antes de impuestos (incluye inversión previa)", "Sin coste de oportunidad; escenarios hipotéticos")
    guardar(fig, "g10_caja_acumulada.png")

    # g11 economía
    e = economia[economia["clase"] != "D"]
    fig, ax = plt.subplots(figsize=(8, 4.2))
    x = np.arange(len(e))
    ax.bar(x - 0.15, e["horas_acumuladas_gratis"], width=0.28, color=SERIES[0], label="Jugador gratuito")
    ax.bar(x + 0.15, e["horas_acumuladas_doble_ganancias"], width=0.28, color=SERIES[1], label="Con pase Doble ganancias")
    for i, (a, b) in enumerate(zip(e["horas_acumuladas_gratis"], e["horas_acumuladas_doble_ganancias"])):
        ax.text(i - 0.15, a + 1, f"{a:.0f} h", ha="center", color=TINTA2, fontsize=8)
        ax.text(i + 0.15, b + 1, f"{b:.0f} h", ha="center", color=TINTA2, fontsize=8)
    ax.set_xticks(x)
    ax.set_xticklabels([f"Primer coche clase {c}" for c in e["clase"]], fontsize=8)
    ax.set_ylabel("Horas de juego acumuladas")
    ax.legend(frameon=False, fontsize=8, labelcolor=TINTA2)
    estilo(ax, "Tiempo objetivo para el primer coche de cada clase", "Diseño propuesto; se valida con telemetría de economía en la beta")
    guardar(fig, "g11_economia_progresion.png")


# ---------------------------------------------------------------------------
# 10. Competidores (instantáneas observadas; ver 02_competidores.md)
# ---------------------------------------------------------------------------
COMPETIDORES = pd.DataFrame([
    # juego, ccu_instantanea, visitas, pico_historico, nota de fecha
    ("Driving Empire", 29001, 3.368e9, 89566, "Rolimons ~18-10-2026"),
    ("Emergency Hamburg", 16460, 1.22e9, 84100, "RoWatcher, sin fecha"),
    ("A Dusty Trip", 9375, 2.5e9, None, "Rotrends, jul-ago 2026"),
    ("Car Driving Indonesia", 6700, 4.26e8, None, "Earnaldo, jul-sep 2026 (6,7-13 k)"),
    ("Jailbreak", 10000, 7.989e9, 173000, "RoWatcher/AllThingsHow, jul-ago 2026"),
    ("Taxi Boss", 8555, 8.42e8, 20700, "Rolimons, sin fecha"),
    ("WRONG WAY (rally coop)", 7000, 4.1e6, 14000, "RoWatcher, oct-2026"),
    ("Car Crushers 2", 5679, 1.4e9, 22900, "Rotrends, sin fecha"),
    ("Midnight Chasers", 5455, 3.088e8, 16800, "Rotrends, sep-2026"),
    ("Greenville", 4442, 1.411e9, 34133, "Rolimons, sin fecha"),
    ("Vehicle Legends", 2910, 1.9e9, 29700, "Rolimons/tracker, jul-2026"),
    ("Car Dealership Tycoon", 2891, 2.51e9, 66726, "Rolimons, sin fecha"),
    ("Rally Panic! (coop)", 1773, 1.7e6, None, "Earnaldo, 12-10-2026"),
    ("Southwest Florida", 1180, 7.793e8, 29500, "Rotrends, ≤ago-2026"),
    ("Drive World", 1063, 4.99e8, 24600, "Rotrends, ≤ago-2026"),
    ("Midnight Racing: Tokyo", 826, None, 3000, "RoWatcher, 2026"),
    ("Project Trackday (sim)", 32, 4.52e7, 408, "RoWatcher, 2026"),
], columns=["juego", "ccu_instantanea", "visitas", "pico_historico", "fuente_fecha"])


# ---------------------------------------------------------------------------
# 11. Excel editable con fórmulas
# ---------------------------------------------------------------------------
def nucleo_mensual(anclas, meses: int = 24) -> list[float]:
    """DAU medio en el mes a (a=0..) generado por 1 usuario nuevo/día durante el mes 0."""
    ret = curva_retencion(*anclas, dias=int(DIAS_MES * (meses + 1)) + 1)
    dias_total = int(DIAS_MES * meses) + 31
    entrada = np.zeros(dias_total)
    entrada[:30] = 1.0
    dau = np.array([np.dot(entrada[: t + 1][::-1], ret[: t + 1]) for t in range(dias_total)])
    return [float(dau[int(a * DIAS_MES): int((a + 1) * DIAS_MES)].mean()) for a in range(meses)]


def excel(fin, objetivos, mensual, matriz_p, economia, competidores):
    wb = Workbook()
    negrita = Font(bold=True)
    entrada = PatternFill("solid", fgColor="FFF4CC")
    titulo = Font(bold=True, size=13)

    # --- Supuestos
    ws = wb.active
    ws.title = "Supuestos"
    ws["A1"] = "Supuestos del modelo (celdas amarillas = editables). Fecha de referencia: 22-10-2026"
    ws["A1"].font = titulo
    supuestos = [
        ("cuota_creador", CUOTA_CREADOR, "Parte de los Robux gastados que recibe el creador (pases y productos)"),
        ("devex_estandar", DEVEX_ESTANDAR, "USD por Earned Robux (desde 05-09-2025)"),
        ("devex_18us", DEVEX_18_US, "USD por Earned Robux elegible de EE. UU. 18+ (desde 08-06-2026)"),
        ("eurusd", EURUSD, "USD por EUR, BCE 20-10-2026"),
        ("recompensa_diaria", RECOMPENSA_DIARIA, "R$ por usuario cualificado/día (Creator Rewards)"),
        ("dias_mes", DIAS_MES, "Días por mes"),
        ("sesiones_dia", PERFILES_ENGAGEMENT["base"][0], "HIPÓTESIS: sesiones por DAU y día"),
        ("minutos_sesion", PERFILES_ENGAGEMENT["base"][1], "HIPÓTESIS: minutos por sesión"),
        ("conversion", MONETIZACION["media"].conversion, "HIPÓTESIS: % de DAU que paga en el día (monetización media)"),
        ("arppu_rs", MONETIZACION["media"].arppu_rs, "HIPÓTESIS: R$ brutos por pagador y día"),
        ("cualificados_cr", MONETIZACION["media"].cualificados_cr, "HIPÓTESIS: % de DAU que cuenta para Creator Rewards"),
        ("cuota_18us", MONETIZACION["media"].cuota_18us, "HIPÓTESIS: % de Earned Robux de compras a la tasa 18+ EE. UU."),
        ("coste_herramientas", COSTE_HERRAMIENTAS_MES, "HIPÓTESIS: €/mes suscripciones IA (2 socios)"),
        ("coste_contenido", COSTE_CONTENIDO_MES, "HIPÓTESIS: €/mes arte y sonido"),
        ("coste_otros", COSTE_OTROS_MES, "HIPÓTESIS: €/mes otros"),
        ("marketing_mes", 100.0, "HIPÓTESIS: €/mes marketing"),
        ("horas_mantenimiento", HORAS_MANTENIMIENTO_MES, "HIPÓTESIS: horas/mes de ambos socios tras el lanzamiento"),
        ("coste_hora", COSTE_HORA_OPORTUNIDAD, "HIPÓTESIS: €/h de coste de oportunidad por socio"),
        ("cpp_usd", CPP_USD["base"], "HIPÓTESIS: USD por partida en Ads Manager"),
    ]
    ws.append([])
    ws.append(["Parámetro", "Valor", "Descripción"])
    for c in ws[3]:
        c.font = negrita
    refs = {}
    for i, (clave, valor, desc) in enumerate(supuestos, start=4):
        ws.cell(row=i, column=1, value=clave)
        celda = ws.cell(row=i, column=2, value=valor)
        celda.fill = entrada
        ws.cell(row=i, column=3, value=desc)
        refs[clave] = f"Supuestos!$B${i}"
    fila = len(supuestos) + 5
    ws.cell(row=fila, column=1, value="Derivados").font = negrita
    ws.cell(row=fila + 1, column=1, value="dau_por_ccu")
    ws.cell(row=fila + 1, column=2, value=f"=1440/({refs['sesiones_dia']}*{refs['minutos_sesion']})")
    refs["dau_por_ccu"] = f"Supuestos!$B${fila + 1}"
    ws.cell(row=fila + 2, column=1, value="arpdau_rs")
    ws.cell(row=fila + 2, column=2, value=f"={refs['conversion']}*{refs['arppu_rs']}")
    refs["arpdau_rs"] = f"Supuestos!$B${fila + 2}"
    ws.cell(row=fila + 3, column=1, value="usd_por_dau_dia")
    ws.cell(row=fila + 3, column=2, value=(
        f"={refs['arpdau_rs']}*{refs['cuota_creador']}*((1-{refs['cuota_18us']})*{refs['devex_estandar']}+{refs['cuota_18us']}*{refs['devex_18us']})"
        f"+{refs['cualificados_cr']}*{refs['recompensa_diaria']}*{refs['devex_estandar']}"))
    refs["usd_por_dau_dia"] = f"Supuestos!$B${fila + 3}"
    ws.cell(row=fila + 4, column=1, value="eur_por_ccu_mes")
    ws.cell(row=fila + 4, column=2, value=f"={refs['dau_por_ccu']}*{refs['usd_por_dau_dia']}*{refs['dias_mes']}/{refs['eurusd']}")
    refs["eur_por_ccu_mes"] = f"Supuestos!$B${fila + 4}"
    ws.cell(row=fila + 5, column=1, value="costes_fijos_mes")
    ws.cell(row=fila + 5, column=2, value=f"={refs['coste_herramientas']}+{refs['coste_contenido']}+{refs['coste_otros']}+{refs['marketing_mes']}")
    refs["costes_fijos_mes"] = f"Supuestos!$B${fila + 5}"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 80

    # --- CCU -> ingresos (fórmulas)
    ws2 = wb.create_sheet("CCU_a_ingresos")
    cab = ["CCU medio", "DAU", "Pagadores/día", "R$ brutos/día", "Earned R$/mes", "USD/mes", "EUR/mes", "EUR/año",
           "Costes EUR/mes", "Beneficio operativo EUR/mes", "Por socio EUR/mes", "Beneficio con coste oportunidad"]
    ws2.append(["Conversión CCU → ingresos con los supuestos de la hoja Supuestos (monetización media editable)"])
    ws2["A1"].font = titulo
    ws2.append(cab)
    for c in ws2[2]:
        c.font = negrita
    for i, ccu in enumerate(OBJETIVOS_CCU, start=3):
        ws2.cell(row=i, column=1, value=ccu).fill = entrada
        ws2.cell(row=i, column=2, value=f"=A{i}*{refs['dau_por_ccu']}")
        ws2.cell(row=i, column=3, value=f"=B{i}*{refs['conversion']}")
        ws2.cell(row=i, column=4, value=f"=B{i}*{refs['arpdau_rs']}")
        ws2.cell(row=i, column=5, value=f"=(D{i}*{refs['cuota_creador']}+B{i}*{refs['cualificados_cr']}*{refs['recompensa_diaria']})*{refs['dias_mes']}")
        ws2.cell(row=i, column=6, value=f"=B{i}*{refs['usd_por_dau_dia']}*{refs['dias_mes']}")
        ws2.cell(row=i, column=7, value=f"=F{i}/{refs['eurusd']}")
        ws2.cell(row=i, column=8, value=f"=G{i}*12")
        ws2.cell(row=i, column=9, value=f"={refs['costes_fijos_mes']}+IF(G{i}>=10000,0.15*G{i},IF(G{i}>=3000,0.1*G{i},0))")
        ws2.cell(row=i, column=10, value=f"=G{i}-I{i}")
        ws2.cell(row=i, column=11, value=f"=J{i}/2")
        ws2.cell(row=i, column=12, value=f"=J{i}-{refs['horas_mantenimiento']}*{refs['coste_hora']}")
    for col in range(1, len(cab) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 16

    # --- Objetivos de ingresos (fórmulas)
    ws3 = wb.create_sheet("Objetivos_ingresos")
    ws3.append(["CCU y DAU necesarios para cada objetivo de ingresos netos mensuales (antes de costes)"])
    ws3["A1"].font = titulo
    ws3.append(["Objetivo EUR/mes", "CCU medio necesario", "DAU necesario", "Pagadores/día", "Nuevos/día (días de vida editables)", "Días de vida por usuario (365 d)"])
    for c in ws3[2]:
        c.font = negrita
    vida_base = dias_vida(curva_retencion(*ESCENARIOS_RETENCION["base"]))
    for i, obj in enumerate(OBJETIVOS_EUR, start=3):
        ws3.cell(row=i, column=1, value=obj).fill = entrada
        ws3.cell(row=i, column=2, value=f"=A{i}/{refs['eur_por_ccu_mes']}")
        ws3.cell(row=i, column=3, value=f"=B{i}*{refs['dau_por_ccu']}")
        ws3.cell(row=i, column=4, value=f"=C{i}*{refs['conversion']}")
        ws3.cell(row=i, column=6, value=round(vida_base, 3)).fill = entrada
        ws3.cell(row=i, column=5, value=f"=C{i}/F{i}")
    for col in range(1, 7):
        ws3.column_dimensions[get_column_letter(col)].width = 22

    # --- Simulador de cohortes mensual (fórmulas)
    ws4 = wb.create_sheet("Simulador_cohortes")
    ws4.append(["Simulador mensual: edita los nuevos usuarios/día de cada mes y el núcleo de retención"])
    ws4["A1"].font = titulo
    ws4.append(["El núcleo K(a) es el DAU medio en el mes a generado por 1 usuario nuevo/día durante un mes (precalculado desde la curva diaria)."])
    nucleos = {k: nucleo_mensual(v) for k, v in ESCENARIOS_RETENCION.items()}
    ws4.append(["Edad (meses)"] + list(range(24)))
    for nombre, nucleo in nucleos.items():
        ws4.append([f"K {nombre}"] + [round(v, 5) for v in nucleo])
    ws4.append(["K usado (editable)"] + [f"={get_column_letter(c)}5" for c in range(2, 26)])  # por defecto: base
    fila_k = ws4.max_row
    for c in range(2, 26):
        ws4.cell(row=fila_k, column=c).fill = entrada
    ws4.append([])
    ws4.append(["Mes"] + list(range(1, 25)))
    fila_mes = ws4.max_row
    base_mensual = mensual["base"]
    ws4.append(["Nuevos usuarios/día (media del mes)"] + [round(v) for v in base_mensual["nuevos_dia_medio"]])
    fila_n = ws4.max_row
    for c in range(2, 26):
        ws4.cell(row=fila_n, column=c).fill = entrada
    ws4.append(["DAU medio"])
    fila_dau = ws4.max_row
    for m in range(24):
        terminos = [f"{get_column_letter(2 + j)}{fila_n}*{get_column_letter(2 + m - j)}{fila_k}" for j in range(m + 1)]
        ws4.cell(row=fila_dau, column=2 + m, value="=" + "+".join(terminos))
    ws4.append(["CCU medio"] + [f"={get_column_letter(c)}{fila_dau}*{refs['sesiones_dia']}*{refs['minutos_sesion']}/1440" for c in range(2, 26)])
    ws4.append(["Ingresos EUR/mes"] + [f"={get_column_letter(c)}{fila_dau}*{refs['usd_por_dau_dia']}*{refs['dias_mes']}/{refs['eurusd']}" for c in range(2, 26)])
    fila_ing = ws4.max_row
    ws4.append(["Costes EUR/mes"] + [f"={refs['costes_fijos_mes']}+IF({get_column_letter(c)}{fila_ing}>=10000,0.15*{get_column_letter(c)}{fila_ing},IF({get_column_letter(c)}{fila_ing}>=3000,0.1*{get_column_letter(c)}{fila_ing},0))" for c in range(2, 26)])
    fila_cos = ws4.max_row
    ws4.append(["Beneficio operativo EUR/mes"] + [f"={get_column_letter(c)}{fila_ing}-{get_column_letter(c)}{fila_cos}" for c in range(2, 26)])
    fila_ben = ws4.max_row
    inversion = sum(INVERSION_PREVIA.values()) + 250
    ws4.append(["Inversión previa EUR", inversion])
    ws4.cell(row=ws4.max_row, column=2).fill = entrada
    fila_inv = ws4.max_row
    ws4.append(["Caja acumulada EUR"] + [f"=SUM($B${fila_ben}:{get_column_letter(c)}{fila_ben})-$B${fila_inv}" for c in range(2, 26)])
    ws4.column_dimensions["A"].width = 36

    # --- Escenarios simulados (valores)
    ws5 = wb.create_sheet("Escenarios_24m")
    ws5.append(["Resultados de la simulación diaria por cohortes (valores generados por modelo.py)"])
    ws5["A1"].font = titulo
    todas = pd.concat(mensual.values())
    ws5.append(list(todas.columns))
    for c in ws5[2]:
        c.font = negrita
    for fila in todas.itertuples(index=False):
        ws5.append([round(v, 2) if isinstance(v, float) else v for v in fila])

    # --- Matriz de conceptos (fórmulas)
    ws6 = wb.create_sheet("Matriz_conceptos")
    ws6.append(["Matriz de decisión ponderada (edita pesos y puntuaciones)"])
    ws6["A1"].font = titulo
    ws6.append(["Criterio", "Peso"] + list(matriz_p.index))
    for c in ws6[2]:
        c.font = negrita
    for i, (clave, nombre, peso) in enumerate(CRITERIOS, start=3):
        ws6.cell(row=i, column=1, value=nombre)
        ws6.cell(row=i, column=2, value=peso).fill = entrada
        for j, concepto in enumerate(matriz_p.index, start=3):
            ws6.cell(row=i, column=j, value=int(matriz_p.loc[concepto, clave])).fill = entrada
    ult = 2 + len(CRITERIOS)
    ws6.cell(row=ult + 1, column=1, value="Puntuación 0-100").font = negrita
    for j in range(3, 3 + len(matriz_p.index)):
        col = get_column_letter(j)
        ws6.cell(row=ult + 1, column=j, value=f"=SUMPRODUCT($B$3:$B${ult},{col}3:{col}{ult})/SUM($B$3:$B${ult})*10")
    ws6.column_dimensions["A"].width = 44
    for j in range(3, 3 + len(matriz_p.index)):
        ws6.column_dimensions[get_column_letter(j)].width = 22
        ws6.cell(row=2, column=j).alignment = Alignment(wrap_text=True)

    # --- Economía (fórmulas)
    ws7 = wb.create_sheet("Economia_juego")
    ws7.append(["Tiempo para el primer coche de cada clase (créditos y horas; edita precios e ingresos)"])
    ws7["A1"].font = titulo
    ws7.append(["Clase", "Precio coche", "Ingreso/hora en clase previa", "Misiones (créditos/h)", "% en mejoras", "Bono inicial", "Horas tramo", "Horas acumuladas"])
    for c in ws7[2]:
        c.font = negrita
    r0 = 3
    for i, (clase, precio, ingreso, mejoras) in enumerate(CLASES[1:], start=r0):
        ws7.cell(row=i, column=1, value=clase)
        for col, v in ((2, precio), (3, ingreso), (4, MISIONES_DIARIAS_HORA), (5, mejoras), (6, BONO_INICIAL if clase == "C" else 0)):
            ws7.cell(row=i, column=col, value=v).fill = entrada
        ws7.cell(row=i, column=7, value=f"=(B{i}-F{i})/((C{i}+D{i})*(1-E{i}))")
        ws7.cell(row=i, column=8, value=f"=SUM($G${r0}:G{i})")
    for col in range(1, 9):
        ws7.column_dimensions[get_column_letter(col)].width = 20

    # --- Competidores (valores)
    ws8 = wb.create_sheet("Competidores")
    ws8.append(list(competidores.columns))
    for fila in competidores.itertuples(index=False):
        ws8.append([None if (isinstance(v, float) and math.isnan(v)) else v for v in fila])
    ws8.column_dimensions["A"].width = 28
    ws8.column_dimensions["E"].width = 40

    wb.save(os.path.join(RAIZ, "modelo_financiero.xlsx"))


# ---------------------------------------------------------------------------
# 12. Autocomprobaciones (aritmética crítica)
# ---------------------------------------------------------------------------
def autocomprobaciones():
    # CCU = DAU × sesiones × minutos / 1440
    assert abs(64.2857 - dau_por_ccu((1.6, 14.0))) < 1e-3
    # 1.000 R$ brutos -> 700 Earned -> 2,66 USD a la tasa estándar
    m = Monetizacion(1.0, 1000, 0.0, 0.0)
    assert abs(ingresos_por_dau_dia(m)["usd"] - 2.66) < 1e-9
    # La curva pasa por las anclas
    c = curva_retencion(0.12, 0.03, 0.01)
    assert abs(c[1] - 0.12) < 1e-9 and abs(c[7] - 0.03) < 1e-9 and abs(c[30] - 0.01) < 1e-9
    # Con entrada constante N y horizonte largo, DAU estacionario ≈ N × días de vida
    e = EscenarioCrecimiento("t", "base", (1.6, 14.0), "media", 0, 1000, 1e9, 1.0, 10_000, 0.0, 0.0, 0, 0, tam_usuarios=1e12)
    df = simular(e, dias=366)
    assert abs(df["dau"].iloc[365] / (1000 * dias_vida(c)) - 1) < 0.01
    # Efecto implícito plataforma: 26,6 % del gasto del jugador llega al creador a un precio medio de ~0,01 USD/R$
    assert abs(CUOTA_CREADOR * DEVEX_ESTANDAR / 0.266 - 0.0100) < 0.0002


if __name__ == "__main__":
    autocomprobaciones()
    semana = perfil_horario_semanal()
    pd.DataFrame({"hora_semana_utc": range(168), "ccu_relativo": semana.round(3)}).to_csv(os.path.join(SALIDAS, "perfil_semanal.csv"), index=False)

    vidas = pd.DataFrame([{"escenario": k, "D1": v[0], "D7": v[1], "D30": v[2],
                           "D90_extrapolado": round(curva_retencion(*v)[90], 4),
                           "dias_activos_365": round(dias_vida(curva_retencion(*v)), 3)} for k, v in ESCENARIOS_RETENCION.items()])
    vidas.to_csv(os.path.join(SALIDAS, "retencion_escenarios.csv"), index=False)

    tabla_c = tabla_ccu(semana)
    tabla_c.to_csv(os.path.join(SALIDAS, "ccu_escenarios.csv"), index=False)

    fin = tabla_finanzas_por_ccu()
    fin.to_csv(os.path.join(SALIDAS, "finanzas_por_ccu.csv"), index=False)

    objetivos = tabla_objetivos()
    objetivos.to_csv(os.path.join(SALIDAS, "objetivos_ingresos.csv"), index=False)

    tabla_ltv().to_csv(os.path.join(SALIDAS, "ltv_y_cpp.csv"), index=False)
    tabla_equilibrio().to_csv(os.path.join(SALIDAS, "punto_equilibrio.csv"), index=False)

    pico_ratio = float(semana.max())
    sims, mensual = {}, {}
    for e in ESCENARIOS:
        sims[e.nombre] = simular(e)
        mensual[e.nombre] = resumen_mensual(e, sims[e.nombre], pico_ratio)
    pd.concat(mensual.values()).round(2).to_csv(os.path.join(SALIDAS, "simulacion_mensual.csv"), index=False)
    hitos = []
    for nombre, m in mensual.items():
        for mes in (3, 6, 12, 24):
            fila = m[m["mes"] == mes].iloc[0]
            acumulado = m[m["mes"] <= mes]
            hitos.append({"escenario": nombre, "mes": mes, "nuevos_dia": round(fila["nuevos_dia_medio"]),
                          "dau_medio": round(fila["dau_medio"]), "ccu_medio": round(fila["ccu_medio"], 1),
                          "ccu_pico_semanal_est": round(fila["ccu_pico_semanal_est"]),
                          "ingresos_eur_mes": round(fila["ingresos_eur"]), "beneficio_eur_mes": round(fila["beneficio_operativo_eur"]),
                          "ingresos_acumulados_eur": round(acumulado["ingresos_eur"].sum()),
                          "caja_acumulada_eur": round(fila["caja_acumulada_eur"])})
    pd.DataFrame(hitos).to_csv(os.path.join(SALIDAS, "simulacion_hitos.csv"), index=False)

    efectos = []
    base_e = ESCENARIOS[1]
    for mkt in (0, 50, 100, 250, 500, 1000):
        df = simular(EscenarioCrecimiento(**{**base_e.__dict__, "marketing_mes_eur": mkt}), dias=366)
        efectos.append({"palanca": "marketing_eur_mes", "valor": mkt,
                        "ccu_mes6": round(df[(df.dia >= 152) & (df.dia < 183)].ccu_medio.mean(), 1),
                        "ccu_mes12": round(df[(df.dia >= 335) & (df.dia < 366)].ccu_medio.mean(), 1),
                        "ingresos_12m_eur": round(df.usd_dia.sum() * USD_A_EUR),
                        "gasto_marketing_12m_eur": round(mkt * 12)})
    for cada in (14, 21, 42, 10_000):
        df = simular(EscenarioCrecimiento(**{**base_e.__dict__, "cada_actualizacion": cada,
                                            "suelo_interes": base_e.suelo_interes if cada <= 42 else 0.15}), dias=366)
        efectos.append({"palanca": "dias_entre_actualizaciones", "valor": cada if cada < 10_000 else "sin actualizaciones",
                        "ccu_mes6": round(df[(df.dia >= 152) & (df.dia < 183)].ccu_medio.mean(), 1),
                        "ccu_mes12": round(df[(df.dia >= 335) & (df.dia < 366)].ccu_medio.mean(), 1),
                        "ingresos_12m_eur": round(df.usd_dia.sum() * USD_A_EUR), "gasto_marketing_12m_eur": None})
    pd.DataFrame(efectos).to_csv(os.path.join(SALIDAS, "efectos_palancas.csv"), index=False)

    sens = sensibilidad(ESCENARIOS[1])
    sens.to_csv(os.path.join(SALIDAS, "sensibilidad_base.csv"), index=False)

    matriz_p, matriz, mc = matriz_decision()
    matriz.to_csv(os.path.join(SALIDAS, "matriz_conceptos.csv"))
    mc.to_csv(os.path.join(SALIDAS, "matriz_montecarlo.csv"))

    economia = tabla_economia()
    economia.to_csv(os.path.join(SALIDAS, "economia_progresion.csv"), index=False)

    COMPETIDORES.to_csv(os.path.join(SALIDAS, "competidores_instantaneas.csv"), index=False)

    plat_ccu = PLAT_HORAS_TRIM / PLAT_DIAS_TRIM / 24
    plataforma = {
        "ccu_medio_plataforma_q2_2026": round(plat_ccu),
        "horas_por_dau_dia": round(PLAT_HORAS_TRIM / PLAT_DIAS_TRIM / PLAT_DAU, 3),
        "bookings_usd_por_dau_dia": round(PLAT_BOOKINGS_TRIM / PLAT_DIAS_TRIM / PLAT_DAU, 4),
        "bookings_usd_por_ccu_medio_mes": round(PLAT_BOOKINGS_TRIM / PLAT_DIAS_TRIM * DIAS_MES / plat_ccu, 2),
        "devex_usd_por_ccu_medio_mes_aprox": round(PLAT_BOOKINGS_TRIM / PLAT_DIAS_TRIM * DIAS_MES / plat_ccu * PLAT_DEVEX_SOBRE_BOOKINGS, 2),
        "competidores_suma_ccu_instantaneas": int(COMPETIDORES["ccu_instantanea"].sum()),
        "cuota_genero_sobre_ccu_medio_plataforma_%": round(100 * COMPETIDORES["ccu_instantanea"].sum() / plat_ccu, 2),
        "eur_por_ccu_mes_baja": round(eur_por_ccu_mes(MONETIZACION["baja"]), 2),
        "eur_por_ccu_mes_media": round(eur_por_ccu_mes(MONETIZACION["media"]), 2),
        "eur_por_ccu_mes_alta": round(eur_por_ccu_mes(MONETIZACION["alta"]), 2),
        "pico_semanal_sobre_media": round(pico_ratio, 2),
        "minimo_sobre_media": round(float(semana.min()), 2),
        "inversion_previa_eur": sum(INVERSION_PREVIA.values()),
    }
    pd.Series(plataforma).to_csv(os.path.join(SALIDAS, "plataforma_derivados.csv"), header=["valor"])

    graficos(semana, sims, mensual, fin, objetivos, sens, matriz, COMPETIDORES, economia)
    excel(fin, objetivos, mensual, matriz_p, economia, COMPETIDORES)
    print("OK", plataforma)
